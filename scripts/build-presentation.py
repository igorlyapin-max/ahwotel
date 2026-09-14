#!/usr/bin/env python3
"""Build the Russian executive deck; requires python-pptx 1.0.2.

All diagrams and text remain editable. Generated illustrations are embedded.
Speaker notes come from docs/PRESENTATION-AHWOTel.md.
"""
from pathlib import Path
import json
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
ASSETS = DOCS / 'presentation-assets'
OUT = DOCS / 'AHWOTel-Project-Overview.pptx'
W, H = 13.333333, 7.5
FONT = 'Noto Sans'
BG, WHITE = 'F3F6FA', 'FFFFFF'
NAVY, INK, MUTED = '102B46', '223D56', '60768A'
BLUE, TEAL, AMBER, RED = '2369DC', '008F95', 'B77814', 'C64F5D'
PALE, AQUA, WARM, GRAY = 'E7EFFD', 'DEF3F0', 'FFF1D9', 'DCE4EB'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'AHWOTel — мониторинг Android: прототип и пилот'
prs.core_properties.subject = 'Производительность, OEM/Knox, Self Telemetry, OpenTelemetry и затраты'
prs.core_properties.author = 'AHWOTel project'
prs.core_properties.keywords = 'AHWOTel, Android, Knox, OpenTelemetry, SOTI'
prs.core_properties.comments = 'Prepared from project specifications and validation reports. Snapshot: 2026-09-14.'


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=True, width=1):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = 0.12
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    if line:
        sh.line.color.rgb = color(line)
        sh.line.width = Pt(width)
    else:
        sh.line.fill.background()
    return sh


def txt(slide, text, x, y, w, h, size=18, bold=False, c=INK, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.01)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(c)
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.line_spacing = 1.12
    return sh


def line(slide, x1, y1, x2, y2, c=GRAY, width=1.5):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color(c)
    sh.line.width = Pt(width)
    return sh


def circle(slide, x, y, d, fill, label=None, size=15, text_color=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    sh.line.fill.background()
    if label is not None:
        tf = sh.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = str(label)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = color(text_color)
    return sh


def pill(slide, text, x, y, w, fill=PALE, c=BLUE):
    rect(slide, x, y, w, .31, fill)
    txt(slide, text, x+.07, y+.028, w-.14, .25, 9.5, True, c, PP_ALIGN.CENTER)


def base(number, section, title, subtitle='', status=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(BG)
    rect(s, .55, .44, .32, .055, TEAL, radius=False)
    txt(s, f'AHWOTel  /  {section.upper()}', .98, .35, 9.3, .29, 10, True, MUTED)
    txt(s, title, .55, .86, 12.2, .72, 28, True, NAVY)
    if subtitle:
        txt(s, subtitle, .57, 1.60, 12.0, .47, 13, False, MUTED)
    if status:
        pill(s, status, 10.0, .33, 2.78)
    line(s, .55, 7.06, 12.78, 7.06, GRAY, .6)
    txt(s, 'AHWOTel  ·  прототип → управляемый пилот', .55, 7.15, 7.9, .21, 8.5, False, MUTED)
    txt(s, '14.09.2026', 10.6, 7.15, 1.45, .21, 8.5, False, MUTED, PP_ALIGN.RIGHT)
    circle(s, 12.35, 7.12, .28, NAVY, f'{number:02}', 8)
    return s


def icon(slide, kind, x, y, scale=.58, c=BLUE):
    """Small editable geometric pictograms, with consistent stroke weight."""
    def seg(a,b,d,e): line(slide,x+a*scale,y+b*scale,x+d*scale,y+e*scale,c,1.8)
    if kind in ('cpu','ram'):
        rect(slide,x+.18*scale,y+.18*scale,.64*scale,.64*scale,WHITE,c,False,1.8)
        for t in (.3,.5,.7):
            seg(t,0,t,.18);seg(t,.82,t,1);seg(0,t,.18,t);seg(.82,t,1,t)
        if kind=='cpu':rect(slide,x+.34*scale,y+.34*scale,.32*scale,.32*scale,PALE,radius=False)
        else:
            for t in (.32,.5,.68):seg(t,.33,t,.67)
    elif kind=='storage':
        for j in range(3):
            rect(slide,x+.1*scale,y+(.08+j*.29)*scale,.8*scale,.23*scale,WHITE,c,False,1.5)
            circle(slide,x+.72*scale,y+(.155+j*.29)*scale,.06*scale,c)
    elif kind=='battery':
        rect(slide,x+.1*scale,y+.23*scale,.75*scale,.53*scale,WHITE,c,True,1.8)
        rect(slide,x+.85*scale,y+.4*scale,.1*scale,.18*scale,c,radius=False)
        rect(slide,x+.2*scale,y+.33*scale,.4*scale,.33*scale,c,radius=False)
    elif kind=='thermal':
        rect(slide,x+.37*scale,y+.05*scale,.25*scale,.67*scale,WHITE,c,True,1.8)
        circle(slide,x+.25*scale,y+.58*scale,.5*scale,c)
        seg(.49,.22,.49,.7)
    elif kind=='network':
        for a,b in ((.05,.65),(.4,.05),(.75,.65)):
            rect(slide,x+a*scale,y+b*scale,.22*scale,.22*scale,WHITE,c,False,1.5)
        seg(.16,.65,.5,.27);seg(.85,.65,.5,.27);seg(.28,.77,.75,.77)
    elif kind=='clock':
        sh=circle(slide,x+.05*scale,y+.05*scale,.9*scale,WHITE)
        sh.line.color.rgb=color(c);sh.line.width=Pt(1.8)
        seg(.5,.5,.5,.25);seg(.5,.5,.71,.63)
    elif kind=='phone':
        rect(slide,x+.23*scale,y+.02*scale,.56*scale,.97*scale,WHITE,c,True,1.8)
        seg(.34,.82,.67,.82)
    elif kind=='shield':
        sh=slide.shapes.add_shape(MSO_SHAPE.PENTAGON,Inches(x+.1*scale),Inches(y+.05*scale),Inches(.8*scale),Inches(.9*scale))
        sh.rotation=180;sh.fill.solid();sh.fill.fore_color.rgb=color(PALE);sh.line.color.rgb=color(c)
        seg(.3,.45,.45,.6);seg(.45,.6,.73,.31)


def arrow(slide,x,y,w=.5,c=BLUE,left=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW if left else MSO_SHAPE.RIGHT_ARROW,
                             Inches(x),Inches(y),Inches(w),Inches(.2))
    sh.fill.solid();sh.fill.fore_color.rgb=color(c);sh.line.fill.background()


def card(slide,x,y,w,h,title,body,kind=None,accent=BLUE):
    rect(slide,x,y,w,h)
    top=y+.22
    if kind:
        icon(slide,kind,x+.22,top,.43,accent)
        txt(slide,title,x+.82,top-.01,w-1.0,.53,16,True,NAVY)
    else:
        rect(slide,x+.22,top,.045,.37,accent,radius=False)
        txt(slide,title,x+.4,top-.01,w-.62,.51,16,True,NAVY)
    compact = h < 1.8
    body_top = .82 if compact else .89
    txt(slide,body,x+.23,y+body_top,w-.46,h-body_top-.12,12 if compact else 13,False,MUTED)


def photo(slide,name,x,y,w,h):
    path=ASSETS/name
    if not path.exists():raise FileNotFoundError(f'Missing presentation illustration: {path}')
    slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))


# 1 — product purpose + conceptual timeline
s=base(1,'Цель проекта','Когда устройство мешает работе',status='ЦЕЛЕВАЯ МЕТОДИКА')
photo(s,'phone.png',6.53,1.75,6.2,3.49)
txt(s,'От технических признаков\nк потерянному рабочему времени',.58,1.95,6.22,.92,23,True,NAVY)
txt(s,'Отдельный Android-агент.\nБизнес-приложение остаётся без изменений.',.6,3.04,5.8,.65,16,False,MUTED)
pill(s,'БИЗНЕС: операция не выполнена',.6,4.00,5.75,WARM,AMBER)
pill(s,'ТЕХНИКА: ресурсы ограничивают работу',.6,4.44,5.75,PALE,BLUE)
rect(s,.55,5.38,12.23,1.46,WHITE)
txt(s,'ИНТЕРВАЛЫ В РАБОЧИЕ ЧАСЫ',.79,5.53,3.22,.25,10,True,MUTED)
for j,(label,segments,col) in enumerate([
    ('Рабочее окно',[(0,1)],PALE),
    ('Техническая деградация',[(.19,.31),(.69,.20)],AMBER),
    ('Подтверждённый простой',[(.26,.18),(.73,.09)],RED),
]):
    yy=5.89+j*.255
    txt(s,label,.8,yy-.02,3.1,.23,10,False,MUTED)
    rect(s,4.0,yy,7.6,.115,BG,radius=False)
    for pos,length in segments:rect(s,4+pos*7.6,yy,length*7.6,.115,col,radius=False)
txt(s,'Схема, не замеры',11.67,5.85,.82,.75,8,False,MUTED)

# 2 — capability grid
s=base(2,'Источники Android','Не каждый показатель доступен любому APK',
       'Sandbox, версия Android, прошивка и права определяют границы измерений.', 'РЕАЛИЗОВАНО В APK')
cards=[('Память','RAM и low memory','Не объясняет причины\nпроблем приложения.','ram'),
       ('Хранилище','Ёмкость и свободное место','Заполненность не равна\nскорости чтения и записи.','storage'),
       ('Батарея','Заряд, температура, контекст','Ток и счётчик заряда\nдоступны не на всех моделях.','battery'),
       ('Thermal','Статус и запас до ограничений','Зависит от версии API\nи реализации устройства.','thermal'),
       ('CPU','CPU агента и косвенные сигналы','Общесистемная загрузка\nне гарантирована.','cpu')]
for i,(title,lead,body,kind) in enumerate(cards):
    x=.55+i*2.49
    rect(s,x,2.35,2.27,3.25,WHITE)
    icon(s,kind,x+.23,2.63,.56,BLUE if i<4 else TEAL)
    txt(s,title,x+.22,3.38,1.88,.4,18,True,NAVY)
    txt(s,lead,x+.22,3.99,1.88,.67,13,True,INK)
    txt(s,body,x+.22,4.84,1.86,.58,10.5,False,MUTED)
rect(s,.55,5.9,12.23,.89,NAVY)
txt(s,'Проверка из APK  →  источник / fallback  →  явный статус',.8,6.07,8.9,.33,18,True,WHITE)
txt(s,'UNSUPPORTED ≠ 0',10.03,6.1,2.43,.36,16,True,'8BE0D6')

# 3 — OEM illustration and support gates
s=base(3,'OEM / Samsung Knox','Больше корпоративного контекста — при наличии прав',
       'Регистрация разработчика и наличие библиотеки сами по себе не открывают все API.', 'KNOX: НУЖЕН СТЕНД')
photo(s,'shield.png',.55,2.16,5.02,2.824)
for i,(a,b) in enumerate([
    ('Устройство и безопасность','Платформа, конфигурация, доступные состояния'),
    ('Приложения и сеть','Инвентарь и ограничения в рамках разрешённых API'),
    ('Корпоративные политики','USB, камера, Wi-Fi, Bluetooth, установка приложений')]):
    yy=2.28+i*1.08
    circle(s,5.88,yy+.06,.33,TEAL,str(i+1),11)
    txt(s,a,6.43,yy,6.1,.43,17,True,NAVY)
    txt(s,b,6.43,yy+.48,6.02,.46,12.5,False,MUTED)
pill(s,'МОДЕЛЬ',.62,5.30,1.52)
pill(s,'ВЕРСИЯ API',2.35,5.30,1.78)
pill(s,'ПРАВА',4.34,5.30,1.46)
pill(s,'ЛИЦЕНЗИЯ*',6.01,5.30,1.84)
pill(s,'ВЫЗОВ ИЗ APK',8.06,5.30,2.08,AQUA,TEAL)
txt(s,'* зависит от API',10.4,5.34,2.04,.25,10,False,MUTED)
rect(s,.55,5.93,12.23,.88,WARM)
txt(s,'SM-J260F: библиотека есть, требуемый публичный метод API отсутствует.',.81,6.06,11.57,.28,14,True,INK)
txt(s,'Расширенные Knox-чтения не подтверждены. Android-сбор продолжает работу.',.81,6.42,11.57,.24,11.5,False,MUTED)

# 4 — self-observability
s=base(4,'Self Telemetry','Измеряем стоимость самого мониторинга',
       'Сравниваем режимы на одинаковых моделях и при сопоставимой рабочей нагрузке.', 'СБОР ЕСТЬ · НУЖЕН A/B')
items=[('CPU и память','Потребление процесса\nи динамика ресурсов','cpu'),('Сеть и хранение','Трафик, база\nи очередь отправки','network'),('Задачи и WakeLock','Длительность, задержки,\nповторы и удержание CPU','clock')]
for i,(title,body,kind) in enumerate(items):card(s,.55+i*4.15,2.26,3.94,1.95,title,body,kind,TEAL if i==2 else BLUE)
txt(s,'ПРОГРАММА СРАВНЕНИЯ',.59,4.53,5,.29,10,True,MUTED)
for i,(a,b) in enumerate([('Агент выключен','Базовый режим'),('Обычный профиль','Штатная нагрузка'),('Диагностика','Повышенная частота')]):
    x=.55+i*4.15
    rect(s,x,4.98,3.94,.92,AQUA if i==1 else WHITE)
    txt(s,a,x+.23,5.13,3.45,.33,15,True,TEAL if i==1 else NAVY)
    txt(s,b,x+.23,5.53,3.45,.25,10.5,False,MUTED)
txt(s,'Self Telemetry < 5% нагрузки агента — целевой бюджет, ещё не доказанный результат.',.63,6.24,12.0,.34,13,True,NAVY)
txt(s,'Разряд батареи устройства — контекст; точное энергопотребление агента в mAh не заявляется.',.63,6.63,12.0,.26,10.5,False,MUTED)

# 5 — editable system architecture
s=base(5,'OpenTelemetry','Готовая инфраструктура от телефона до витрины',
       'OTel — инструменты для телеметрии. OTLP — протокол передачи. Хранение и графики — отдельные системы.', 'ВНЕШНИЙ СТЕНД — ДАЛЕЕ')
nodes=[(.55,2.35,'AHWOTel','Локальная история\nи очередь',2.33),
       (4.12,2.35,'OTel Collector','Защищённый ingest\nПриём → обработка → экспорт',3.02),
       (8.1,2.35,'Prometheus','Хранение метрик',1.95),
       (10.79,2.35,'Grafana','Графики и витрины',1.99)]
for x,y,a,b,w in nodes:
    rect(s,x,y,w,1.86,WHITE)
    txt(s,a,x+.17,y+.24,w-.34,.43,17 if w>2 else 13.5,True,NAVY)
    txt(s,b,x+.17,y+.99,w-.34,.65,11,False,MUTED)
arrow(s,3.03,3.17,.8,BLUE)
txt(s,'PUSH',3.02,2.6,.92,.24,10,True,BLUE,PP_ALIGN.CENTER)
txt(s,'OTLP/HTTPS',2.92,3.66,1.1,.27,8,False,MUTED,PP_ALIGN.CENTER)
arrow(s,7.38,3.17,.47,TEAL,True)
txt(s,'PULL',7.32,2.6,.67,.25,9,True,TEAL,PP_ALIGN.CENTER)
arrow(s,10.22,3.17,.39,MUTED,True)
txt(s,'запрос',10.12,3.62,.64,.25,8,False,MUTED,PP_ALIGN.CENTER)
txt(s,'Стрелки — инициаторы запросов; метрики возвращаются в ответах.',.62,4.42,12,.26,10.5,False,MUTED)
card(s,.55,5.0,3.94,1.75,'Управление','SOTI: установка, настройки, START/STOP',accent=TEAL)
card(s,4.7,5.0,3.94,1.75,'Расширяемость','Receivers · processors · exporters;\nслужебные extensions',accent=BLUE)
card(s,8.85,5.0,3.93,1.75,'Текущий результат','Локальный HTTPS mock — проверено.\nВнешний Collector — далее.',accent=AMBER)

# 6 — security identity
s=base(6,'Доверенная телеметрия','Кто отправляет данные — и за какое устройство?',
       'Защита канала, аутентификация отправителя и авторизация приёма решают разные задачи.', 'MTLS / ENROLLMENT — ДАЛЕЕ')
for i,(label,head,body) in enumerate([
    ('01','HTTPS','Защита канала\nи проверка сервера'),
    ('02','mTLS','Сертификат каждого устройства;\nключ в Android Keystore'),
    ('03','Авторизация','Identity сертификата\nсоответствует device.id')]):
    x=.55+i*4.15
    rect(s,x,2.3,3.94,2.34,WHITE)
    circle(s,x+.23,2.57,.43,BLUE if i==0 else TEAL,label,12)
    txt(s,head,x+.83,2.59,2.86,.4,19,True,NAVY)
    txt(s,body,x+.24,3.46,3.45,.79,15,False,MUTED)
txt(s,'ЖИЗНЕННЫЙ ЦИКЛ ДОСТУПА',.62,4.97,7,.29,10,True,MUTED)
steps=['Регистрация','Выдача','Обновление','Отзыв']
for i,step in enumerate(steps):
    x=.55+i*3.15
    pill(s,step,x,5.45,2.67,AQUA,TEAL)
    if i<3:arrow(s,x+2.79,5.50,.24,TEAL)
rect(s,.55,6.10,12.23,.69,NAVY)
txt(s,'Сейчас: HTTPS. Цель: индивидуальный доступ. Альтернатива mTLS: короткоживущий JWT.',.8,6.3,11.73,.29,12.2,True,WHITE)

# 7 — weekend timeline
s=base(7,'Vibe coding','От обсуждения до работающего APK за выходные',
       '11–13 сентября 2026. Хронология автора, три коммита Git и отдельные испытания на Samsung.')
for i,(a,b) in enumerate([('ПТ · 11 СЕН','Консультации: актуальность\nи варианты приёма данных'),('СБ · 12 СЕН','Обдумывание идеи:\n«делать или не делать?»')]):
    x=.55+i*6.24
    rect(s,x,2.21,5.99,.94,WHITE)
    txt(s,a,x+.2,2.38,1.54,.25,10.5,True,TEAL)
    txt(s,b,x+1.87,2.36,3.84,.65,12,False,MUTED)
line(s,.96,3.71,12.19,3.71,GRAY,3)
events=[('07:30','Три блока ТЗ','Ресурсы · OEM/Knox\nSelf Telemetry','9bfecfb · LICENSE'),
        ('14:01','Автономный MVP','Метрики · история · графики\nПодсказки EN/RU','5a8710d'),
        ('20:42','OEM + Battery/Self','Новые настройки и оси\nИсправления по review','beaf486'),
        ('ВЕЧЕР','Испытания Samsung','Установка и новый сбор\n20 native-сценариев','После последнего коммита')]
for i,(time,title,body,commit) in enumerate(events):
    x=.55+i*3.12
    circle(s,x+.19,3.49,.43,TEAL if i==3 else BLUE,str(i+1),12)
    txt(s,time,x+.06,4.05,2.85,.47,26 if i<3 else 22,True,NAVY)
    txt(s,title,x+.06,4.69,2.87,.38,15.5,True,INK)
    txt(s,body,x+.06,5.26,2.84,.69,11.3,False,MUTED)
    pill(s,commit,x+.02,6.1,2.90,PALE,BLUE)
txt(s,'Москва, UTC+03:00. Подготовка ТЗ — со слов автора; Git фиксирует время коммитов, а не длительность работы.',.59,6.67,12.13,.26,9,False,MUTED)

# 8 — next steps roadmap
s=base(8,'Следующие шаги','Управляемый пилот с общей витриной',
       'Проверяем весь путь: команда SOTI → измерения → защищённый приём → графики.', 'НУЖНО СОГЛАСОВАТЬ ПИЛОТ')
stages=[('Методика и доступы','Бизнес-критерии и календарь\nKnox Developer и лицензии\nGoogle developer: по маршруту поставки'),
        ('Устройство под SOTI','Поддерживаемый Samsung\nУстановка и START/STOP\nПроверка прав самого APK'),
        ('Стенд и витрина','Защищённый ingest и PKI\nOTel Collector + Prometheus\nGrafana: ресурсы и стоимость агента'),
        ('Приёмка пилота','Offline/reconnect и timeout\nScreen-off и нагрузка агента\nВремя, retention и восстановление')]
for i,(a,b) in enumerate(stages):
    x=.55+i*3.12
    rect(s,x,2.37,2.87,3.36,WHITE)
    circle(s,x+.2,2.61,.43,BLUE if i<3 else TEAL,f'{i+1:02}',12)
    txt(s,a,x+.2,3.25,2.45,.82,18,True,NAVY)
    txt(s,b,x+.2,4.37,2.46,1.12,11.1,False,MUTED)
    if i<3:arrow(s,x+2.91,3.66,.16,MUTED)
rect(s,.55,6.08,12.23,.75,AQUA)
txt(s,'Нужны: устройство и SOTI  ·  владелец стенда / PKI  ·  бизнес-эксперт',.81,6.31,11.72,.32,16,True,TEAL)

# 9 — measured usage, author estimate, conditional financial allocation
usage=json.loads((DOCS/'PRESENTATION-USAGE-2026-09-14.json').read_text())
agg=usage['aggregate']
assert agg['total_tokens']==200415729
s=base(9,'Затраты прототипа','Вычисления, решения человека и доля подписки',
       'Снимок 14.09.2026, 08:17 МСК. Включает подготовку презентации до снимка; создание PPTX не включено.')
for i,(value,label,body,accent) in enumerate([
    ('200,4 млн','токенов','188,9 млн — основная сессия\n11,5 млн — семь агентов review',BLUE),
    ('1–2 часа','участия человека','ТЗ, планы/review, согласования\nи проверки телефона — оценка автора',TEAL),
    ('≈ $23–25','условной доли подписки','50% недельного лимита\nпри подписке $200 в месяц',BLUE),
]):
    x=.55+i*4.15
    rect(s,x,2.30,3.94,2.59,WHITE)
    txt(s,value,x+.24,2.55,3.48,.64,35,True,accent)
    txt(s,label,x+.24,3.4,3.45,.44,16,True,NAVY)
    txt(s,body,x+.24,4.06,3.43,.61,11.5,False,MUTED)
txt(s,'СОСТАВ ЗАРЕГИСТРИРОВАННЫХ ТОКЕНОВ',.61,5.24,8,.27,10,True,MUTED)
total=agg['total_tokens'];x=.6
parts=[(agg['cached_input_tokens'],BLUE),(usage['uncached_input_tokens'],TEAL),(agg['output_tokens'],AMBER)]
for count,col in parts:
    width=count/total*12.1
    rect(s,x,5.79,width,.23,col,radius=False);x+=width
for x,col,lab in [(.6,BLUE,'193,1 млн · вход из кэша'),(5.25,TEAL,'6,31 млн · остальной вход'),(9.94,AMBER,'1,01 млн · выход')]:
    circle(s,x,6.24,.13,col)
    txt(s,lab,x+.23,6.20,4.4 if x<9 else 2.4,.31,10.5,False,MUTED)
txt(s,'96,8% входа — кэш. Деньги — распределение фиксированной платы, не отдельное списание и не API-счёт.',.61,6.68,12.1,.23,9.4,False,MUTED)


# Full notes are intentionally outside the visual canvas.
source=(DOCS/'PRESENTATION-AHWOTel.md').read_text()
blocks=re.split(r'^## Слайд \d+\. ',source,flags=re.M)[1:]
assert len(blocks)==len(prs.slides)==9
for i,(slide,block) in enumerate(zip(prs.slides,blocks),1):
    block=block.split('## Редакторская проверка',1)[0]
    notes=block.split('### Заметки докладчика',1)[1].strip()
    # Relative repository references remain readable in PowerPoint's notes pane.
    notes=re.sub(r'\[([^]]+)\]\((?!https?://)([^)]+)\)',lambda m:f'{m[1]} [проект: docs/{m[2]}]',notes)
    notes=re.sub(r'\[([^]]+)\]\((https?://[^)]+)\)',lambda m:f'{m[1]} — {m[2]}',notes)
    notes=notes.replace('**','').replace('`','')
    preface=f'AHWOTel · Слайд {i} · Материалы от 2026-09-14\nИсточник содержания: docs/PRESENTATION-AHWOTel.md\n\n'
    if i in (1,3):
        notes+='\n\nИллюстрация сгенерирована для презентации: концептуальный образ, не фотография тестового телефона и не снимок интерфейса APK.'
    if i==9:
        notes+='\n\nЗатраты зафиксированы до создания PPTX и генерации иллюстраций; эти последующие операции в снимок не включены.'
    slide.notes_slide.notes_text_frame.text=preface+notes

# Catch accidental objects outside the canvas before writing.
for i,slide in enumerate(prs.slides,1):
    for sh in slide.shapes:
        # The default Office shape style inherits a theme shadow.
        # Keep diagrams flat; the illustrations already provide depth.
        style = sh._element.find(qn('p:style'))
        if style is not None:
            effect = style.find(qn('a:effectRef'))
            if effect is not None:
                effect.set('idx', '0')
        assert sh.left>=0 and sh.top>=0,(i,sh.name,'negative position')
        assert sh.left+sh.width<=prs.slide_width+2,(i,sh.name,'outside right edge')
        assert sh.top+sh.height<=prs.slide_height+2,(i,sh.name,'outside bottom edge')
prs.save(OUT)
print(f'Saved {OUT}: {len(prs.slides)} slides, speaker notes, editable shapes and embedded images')
